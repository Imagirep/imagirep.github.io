"""Build the animated teaser (HTML fragment) from the teaser PowerPoint.

Every PowerPoint shape becomes an absolutely positioned element in a 720 x 405 pt coordinate box that scales with the
page (container query units). Fixes relative to the paper figure:
  * The Top-down card answers the obstacle-direction question (Q3) and the Depth card answers the farthest-object
    question (Q2). Their labels, colours, scene boxes and arrows are corrected accordingly.
  * The static dollhouse image is replaced by a rendered camera move (video) that lands on a still of the same view;
    scene boxes are projected from Matterport3D object boxes (tools/render/boxes_canonical.json).
Output: partials/teaser.html (inserted into index.html by tools/build_site.py).
"""
import html, json, os, re
from pptx import Presentation
from lxml import etree

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PPTX = os.environ.get('TEASER_PPTX', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'teaser.pptx'))
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
THEME = {'accent3': '78909C', 'accent6': 'EEFF41', 'lt2': 'EEEEEE', 'dk2': '595959', 'lt1': 'FFFFFF', 'dk1': '000000'}
EMU = 12700.0

# ---- semantic groups and animation steps (shape id -> (step, group))
STEP = {}
def put(step, ids, grp=''):
    for i in ids:
        STEP[i] = (step, grp)
put(1, [60, 72]);                    put(2, [99, 78])
put(3, [73, 55], 'q1'); put(3, [132, 74, 57], 'q2'); put(3, [133, 75, 58], 'q3'); put(3, [134, 76, 56], 'q4'); put(3, [130])
put(4, [59, 77, 121, 85]); put(4, [61], 'q1'); put(4, [62], 'q2'); put(4, [63], 'q3'); put(4, [64], 'q4')
put(5, [65, 86]); put(5.2, [84, 87]); put(5.4, [119, 89, 88, 120])
put(6, [131]); put(6.2, [122, 126], 'q1'); put(6.3, [123, 127], 'q2'); put(6.4, [124, 128], 'q3'); put(6.5, [125, 129], 'q4'); put(6.7, [118])
put(7, [54, 139]); put(7.2, [81, 102, 82, 83])
put(8.0, [66, 98, 71, 100, 79, 135], 'q1')
put(8.3, [68, 80, 106, 115, 105, 136], 'q3')
put(8.6, [69, 101, 110, 111, 112, 114, 104, 137], 'q2')
put(8.9, [70, 113, 107, 108, 116, 109, 138], 'q4')
put(9.4, [103, 117])
SKIP = {67, 90, 92, 94, 96, 91, 93, 95, 97}           # replaced by the video / SVG overlay

prs = Presentation(PPTX)
shapes = {int(s.name.split(';')[1]): s for s in prs.slides[0].shapes}


def color_of(el):
    if el is None:
        return None
    s = el.find('a:srgbClr', NS)
    if s is not None:
        return '#' + s.get('val')
    s = el.find('a:schemeClr', NS)
    if s is not None:
        return '#' + THEME.get(s.get('val'), '000000')
    return None


def pct(v, total):
    return f'{100 * v / total:.4f}%'


def style_box(s):
    x, y, w, h = s.left / EMU, s.top / EMU, s.width / EMU, s.height / EMU
    st = f'left:{pct(x, 720)};top:{pct(y, 405)};width:{pct(w, 720)};height:{pct(h, 405)};'
    if s.rotation:
        st += f'transform:rotate({s.rotation:.2f}deg);'
    return st


def run_html(r):
    rp = r.find('a:rPr', NS)
    t = html.escape(r.findtext('a:t', default='', namespaces=NS))
    css = []
    if rp is not None:
        if rp.get('sz'):
            css.append(f'font-size:calc(var(--u)*{int(rp.get("sz")) / 100:g})')
        if rp.get('b') == '1':
            css.append('font-weight:700')
        if rp.get('i') == '1':
            css.append('font-style:italic')
        c = color_of(rp.find('a:solidFill', NS))
        if c:
            css.append(f'color:{c}')
        lat = rp.find('a:latin', NS)
        if lat is not None and lat.get('typeface'):
            css.append('font-family:' + ('var(--tz-hand)' if 'Comic' in lat.get('typeface') else 'var(--tz-sans)'))
    return f'<span style="{";".join(css)}">{t}</span>' if css else t


def text_html(s):
    e = s._element
    bp = e.find('.//a:bodyPr', NS)
    anchor = {'t': 'flex-start', 'ctr': 'center', 'b': 'flex-end'}.get(bp.get('anchor') if bp is not None else 'ctr', 'center')
    paras = []
    for p in e.findall('.//p:txBody/a:p', NS):
        algn = (p.find('a:pPr', NS).get('algn') if p.find('a:pPr', NS) is not None else None) or 'l'
        inner = ''.join(run_html(r) if r.tag.endswith('}r') else '<br>' for r in p if r.tag.endswith('}r') or r.tag.endswith('}br'))
        paras.append(f'<p style="text-align:{ {"l": "left", "ctr": "center", "r": "right"}.get(algn, "left")}">{inner}</p>' if inner.strip() else '')
    while paras and not paras[0]:
        paras.pop(0)
    while paras and not paras[-1]:
        paras.pop()
    return anchor, ''.join(p or '<p class="tz-gap"></p>' for p in paras)


FIX_TEXT = {105: ('Q2.', 'Q3.'), 104: ('Q3.', 'Q2.')}
FIX_FILL = {68: '#E2EAF0', 69: '#E4ECE3'}            # Top-down -> Q3 blue, Depth -> Q2 green
SWAP_TEXT_COLOR = {80: 101, 101: 80, 105: 104, 104: 105}
orig_colors = {}
for sid in SWAP_TEXT_COLOR:
    cs = [color_of(rp.find('a:solidFill', NS)) for rp in shapes[sid]._element.findall('.//a:rPr', NS)]
    orig_colors[sid] = [c for c in cs if c]


def build_shape(sid, s):
    step, grp = STEP.get(sid, (0, ''))
    attrs = f'class="tz tz-{sid}" data-step="{step}"' + (f' data-q="{grp}"' if grp else '')
    st = style_box(s)
    kind = int(s.shape_type)
    if kind == 13:                                    # picture
        src = f'static/teaser/{sid}.webp'
        alt = 'current view' if sid in (60, 81) else ''
        l, t, r, b = s.crop_left, s.crop_top, s.crop_right, s.crop_bottom
        if any((l, t, r, b)):
            w, h = 1 - l - r, 1 - t - b
            img = (f'<img src="{src}" alt="{alt}" decoding="async" style="position:absolute;width:{100 / w:.3f}%;'
                   f'height:{100 / h:.3f}%;left:{-100 * l / w:.3f}%;top:{-100 * t / h:.3f}%;max-width:none">')
            return f'<div {attrs} style="{st}overflow:hidden">{img}</div>'
        return f'<img {attrs} style="{st}" src="{src}" alt="{alt}" decoding="async">'
    e = s._element
    sp = e.find('p:spPr', NS)
    fill = color_of(sp.find('a:solidFill', NS)) if sp is not None else None
    fill = FIX_FILL.get(sid, fill)
    ln = sp.find('a:ln', NS) if sp is not None else None
    line = color_of(ln.find('a:solidFill', NS)) if ln is not None else None
    lw = int(ln.get('w')) / EMU if ln is not None and ln.get('w') else 0
    geom = sp.find('a:prstGeom', NS).get('prst') if sp is not None and sp.find('a:prstGeom', NS) is not None else 'rect'
    adj = [int(g.get('fmla').split()[1]) for g in sp.findall('a:prstGeom/a:avLst/a:gd', NS)] if sp is not None else []
    w, h = s.width / EMU, s.height / EMU
    css = st
    inner = ''
    if geom == 'roundRect':
        r = (adj[0] if adj else 16667) / 100000 * min(w, h)
        css += f'border-radius:calc(var(--u)*{r:.2f});'
    if geom in ('downArrow', 'mathPlus'):
        col = fill or line or '#595959'
        if geom == 'downArrow':
            inner = (f'<svg viewBox="0 0 {w:.1f} {h:.1f}" preserveAspectRatio="none"><path d="M{w * .25:.1f},0 H{w * .75:.1f} '
                     f'V{h * .5:.1f} H{w:.1f} L{w / 2:.1f},{h:.1f} L0,{h * .5:.1f} H{w * .25:.1f} Z" fill="{col}"/></svg>')
        else:
            t = w * 0.12
            inner = (f'<svg viewBox="0 0 {w:.1f} {h:.1f}"><path d="M{w / 2 - t:.1f},{h * .18:.1f} h{2 * t:.1f} v{h * .32 - t:.1f} '
                     f'h{w * .32 - t:.1f} v{2 * t:.1f} h-{w * .32 - t:.1f} v{h * .32 - t:.1f} h-{2 * t:.1f} v-{h * .32 - t:.1f} '
                     f'h-{w * .32 - t:.1f} v-{2 * t:.1f} h{w * .32 - t:.1f} Z" fill="#595959"/></svg>')
        return f'<div {attrs} style="{css}">{inner}</div>'
    if fill:
        css += f'background:{fill};'
    if line and lw:
        css += f'border:calc(var(--u)*{lw:.2f}) solid {line};'
    if s.has_text_frame and s.text_frame.text.strip():
        anchor, body = text_html(s)
        if sid in FIX_TEXT:
            a, b = FIX_TEXT[sid]; body = body.replace(a, b, 1)
        if sid in SWAP_TEXT_COLOR:
            mine, theirs = orig_colors[sid], orig_colors[SWAP_TEXT_COLOR[sid]]
            for c0, c1 in zip(dict.fromkeys(mine), dict.fromkeys(theirs)):
                body = body.replace(f'color:{c0}', f'color:__{c1}')
            body = body.replace('color:__', 'color:')
        css += f'justify-content:{anchor};'
        return f'<div {attrs} style="{css}"><div class="tz-txt">{body}</div></div>'
    return f'<div {attrs} style="{css}"></div>'


# ---- scene (video + still) placement and projected boxes, in teaser pt units
import numpy as np
from PIL import Image
_im = np.asarray(Image.open(os.path.join(ROOT, 'static', 'teaser', 'scene_canonical.webp')).convert('L')) < 245
_ys, _xs = np.where(_im)
FH, FW = _im.shape
bx0, by0, bx1, by1 = _xs.min(), _ys.min(), _xs.max(), _ys.max()      # rendered scene content (px)
TARGET = (354.0, 168.0, 563.0, 346.0)                                # where the scene sits in the paper figure (pt)
S_PX = min((TARGET[2] - TARGET[0]) / (bx1 - bx0), (TARGET[3] - TARGET[1]) / (by1 - by0))
sx0 = (TARGET[0] + TARGET[2]) / 2 - (bx0 + bx1) / 2 * S_PX
sy0 = (TARGET[1] + TARGET[3]) / 2 - (by0 + by1) / 2 * S_PX
sw, sh = FW * S_PX, FH * S_PX
boxes = json.load(open(os.path.join(ROOT, 'tools', 'render', 'boxes_canonical.json')))
boxes['pillar'] = dict(poly=[[0.538, 0.622], [0.571, 0.622], [0.571, 0.728], [0.538, 0.728]])   # not annotated in MP3D
BOX = {'pillar': ('q1', '#F4CCCC', '#D9BFB8', 66, 'right'), 'table': ('q3', '#6D9EEB', '#78909C', 68, 'left'),
       'sofa': ('q2', '#93C47D', '#5E7860', 69, 'right'), 'pillows': ('q4', '#FF00FF', '#766A82', 70, 'left')}
STEP_OF = {'q1': 8.0, 'q3': 8.3, 'q2': 8.6, 'q4': 8.9}


def to_pt(fx, fy):
    return sx0 + fx * sw, sy0 + fy * sh


svg = []
for name, (q, box_col, arrow_col, card, side) in BOX.items():
    pts = [to_pt(x, y) for x, y in boxes[name]['poly']]
    d = 'M' + ' L'.join(f'{x:.2f},{y:.2f}' for x, y in pts) + ' Z'
    cs = shapes[card]; cx, cy, cw, ch = cs.left / EMU, cs.top / EMU, cs.width / EMU, cs.height / EMU
    bx = sum(p[0] for p in pts) / 4; by = sum(p[1] for p in pts) / 4
    ex, ey = (cx + cw + 1.5, cy + ch * 0.48) if side == 'right' else (cx - 1.5, cy + ch * 0.48)
    # start on the box edge facing the card
    sxp = min(p[0] for p in pts) if side == 'right' else max(p[0] for p in pts)
    path = f'M{sxp:.2f},{by:.2f} C{(sxp + ex) / 2:.2f},{by:.2f} {(sxp + ex) / 2:.2f},{ey:.2f} {ex:.2f},{ey:.2f}'
    st = STEP_OF[q]
    svg.append(f'<path class="tz-box" data-step="{st - 0.15:.2f}" data-q="{q}" d="{d}" stroke="{box_col}"/>')
    svg.append(f'<path class="tz-arrow" data-step="{st - 0.1:.2f}" data-q="{q}" d="{path}" stroke="{arrow_col}" '
               f'marker-end="url(#tz-head-{q})"/>')
    svg.append(f'<defs><marker id="tz-head-{q}" viewBox="0 0 10 10" refX="8.5" refY="5" markerWidth="5" markerHeight="5" '
               f'orient="auto-start-reverse"><path d="M0,0 L10,5 L0,10 z" fill="{arrow_col}"/></marker></defs>')

parts = []
for sid in sorted(shapes, key=lambda i: list(shapes).index(i)):
    if sid in SKIP:
        continue
    parts.append(build_shape(sid, shapes[sid]))
    if sid == 54:   # right panel frame: the scene goes right above it
        scene = (f'<div class="tz tz-scene" data-step="7.2" style="left:{pct(sx0, 720)};top:{pct(sy0, 405)};'
                 f'width:{pct(sw, 720)};height:{pct(sh, 405)}">'
                 '<video class="tz-scene-video" muted playsinline preload="auto" poster="static/teaser/scene_poster.webp">'
                 '<source src="static/video/scene_sweep.webm" type="video/webm">'
                 '<source src="static/video/scene_sweep.mp4" type="video/mp4"></video>'
                 '<img class="tz-scene-still" src="static/teaser/scene_canonical.webp" alt="Imagined 3D state of the scene"></div>')
        parts.append(scene)

frag = ('<div class="teaser" id="teaser" role="img" aria-label="Teaser: the same imagined 3D state, shown through '
        'different representations, answers different spatial questions; generating future RGB views alone does not.">\n'
        + '\n'.join(parts)
        + '\n<svg class="tz-overlay" viewBox="0 0 720 405" aria-hidden="true">' + ''.join(svg) + '</svg>\n'
        '<button class="tz-replay" type="button" aria-label="Replay the teaser animation">Replay</button>\n</div>')
os.makedirs(os.path.join(ROOT, 'partials'), exist_ok=True)
open(os.path.join(ROOT, 'partials', 'teaser.html'), 'w', encoding='utf-8').write(frag)
print('teaser fragment:', len(parts), 'elements,', len(svg) // 3, 'scene boxes; scene box at',
      tuple(round(v, 1) for v in (sx0, sy0, sw, sh)))
